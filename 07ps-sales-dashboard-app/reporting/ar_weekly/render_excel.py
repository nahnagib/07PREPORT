"""Excel workbooks in Arabic (right-to-left sheets): one per salesperson and currency, and one for the team."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from engine import RepReport
from render_html import ccy_ar, day_ar

HEAD_FILL = PatternFill("solid", fgColor="25313F")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
TOT_FILL = PatternFill("solid", fgColor="EEF1F5")
BODY_FONT = Font(name="Arial")
THIN = Border(bottom=Side(style="thin", color="DDE2E8"))
MONEY = '#,##0.00;[Red]-#,##0.00'
INT = '#,##0'
PCT = '0.0%'
DATE = 'dd/mm/yyyy'
RIGHT = Alignment(horizontal="right", vertical="top", wrap_text=True)


def _sheet(wb: Workbook, title: str):
    ws = wb.create_sheet(title) if wb.worksheets and wb.active.title != "Sheet" else wb.active
    ws.title = title
    ws.sheet_view.rightToLeft = True
    return ws


def _write(ws, df: pd.DataFrame, start_row: int = 1, money=(), pct=(), dates=(), ints=(), title: str | None = None,
           fills: dict[int, str] | None = None, total: bool = False, widths: dict[str, int] | None = None) -> int:
    r = start_row
    if title:
        ws.cell(row=r, column=1, value=title).font = Font(name="Arial", bold=True, size=13)
        r += 1
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=r, column=j, value=col)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 32
    for i, row in enumerate(df.itertuples(index=False), 1):
        last = total and i == len(df)
        for j, (col, v) in enumerate(zip(df.columns, row), 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(row=r + i, column=j, value=v)
            c.border, c.font = THIN, Font(name="Arial", bold=last)
            if last:
                c.fill = TOT_FILL
            if col in money:
                c.number_format = MONEY
            elif col in pct:
                c.number_format = PCT
            elif col in dates:
                c.number_format = DATE
            elif col in ints:
                c.number_format = INT
            else:
                c.alignment = RIGHT
        if fills and (i - 1) in fills:
            ws.cell(row=r + i, column=1).fill = PatternFill("solid", fgColor=fills[i - 1].lstrip("#"))
            ws.cell(row=r + i, column=1).font = Font(name="Arial", bold=True, color="FFFFFF")
    for j, col in enumerate(df.columns, 1):
        letter = get_column_letter(j)
        if widths and col in widths:
            w = widths[col]
        else:
            sample = [len(f"{v:,.2f}" if isinstance(v, float) else str(v)) for v in df[col].head(300)] if len(df) else []
            w = min(50, max([len(str(col)) + 2] + [x + 2 for x in sample]))
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, w)
    return r + len(df) + 2


def _table_sheet(wb: Workbook, title: str, df: pd.DataFrame, **kw) -> None:
    ws = _sheet(wb, title)
    _write(ws, df, **kw)
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"


def rep_workbook(r: RepReport, path: Path, generated: dt.datetime) -> Path:
    k, w, c = r.kpi, r.week, ccy_ar(r.currency)
    gross = k["gross"] or 0
    wb = Workbook()

    # 1. summary: the same two tables as the email, plus target and unapplied credits
    ws = _sheet(wb, "الملخص")
    ws["A1"] = "تقرير التحصيل الأسبوعي"
    ws["A1"].font = Font(name="Arial", bold=True, size=15)
    ws["A2"] = f"المندوب: {r.rep}" + (f" — {r.team}" if r.team else "")
    ws["A3"] = f"الفترة: من {day_ar(w.start_date + dt.timedelta(days=1))} إلى {day_ar(w.end_date)} — العملة: {c}"
    ws["A4"] = f"تم الإعداد من نظام Odoo بتاريخ {generated:%d/%m/%Y %H:%M}"
    for a in ("A2", "A3", "A4"):
        ws[a].font = Font(name="Arial", color="5F6B7A")
    status = pd.DataFrame([
        ["إجمالي الديون", k["gross"], 1.0 if gross else None, k["invoices"], k["debtors"]],
        ["المتأخرات (تجاوزت تاريخ الاستحقاق)", k["overdue"], k["overdue"] / gross if gross else None, k["overdue_invoices"], k["overdue_clients"]],
        ["غير مستحقة بعد", k["not_due"], k["not_due"] / gross if gross else None, k["not_due_invoices"], k["not_due_clients"]],
    ], columns=["البند", "المبلغ", "النسبة", "عدد الفواتير", "عدد العملاء"])
    nxt = _write(ws, status, start_row=6, money=("المبلغ",), pct=("النسبة",), ints=("عدد الفواتير", "عدد العملاء"),
                 title=f"الوضع كما في {day_ar(r.as_of)}", widths={"البند": 52, "المبلغ": 18, "النسبة": 12})
    start = w.start_overdue
    rows = [[f"المتأخرات في بداية الأسبوع ({day_ar(w.start_date)})", start, 1.0 if start else None],
            ["المحصّل منها خلال الأسبوع", w.collected_overdue, w.collection_rate]]
    if abs(w.adjustments) >= 0.5:
        rows.append(["إشعارات دائنة وتسويات عليها", w.adjustments, w.adjustments / start if start else None])
    rows += [["المتبقي منها", w.remaining, w.remaining / start if start else None],
             ["متأخرات جديدة دخلت خلال الأسبوع", w.new_overdue, None],
             [f"إجمالي المتأخرات في نهاية الأسبوع ({day_ar(w.end_date)})", w.end_overdue, None]]
    coll = pd.DataFrame(rows, columns=["البند", "المبلغ", "النسبة"])
    nxt = _write(ws, coll, start_row=nxt, money=("المبلغ",), pct=("النسبة",), title="التحصيل الأسبوعي")
    extra = pd.DataFrame([
        ["المستهدف الأسبوعي", w.target, None],
        ["نسبة تحقيق المستهدف", None, w.achievement],
        ["إجمالي التحصيل النقدي خلال الأسبوع (يشمل فواتير غير متأخرة)", w.collected_total, None],
        ["أرصدة دائنة غير مخصصة (دفعات وإشعارات دائنة لم تُطابق مع فواتير)", k["credits"], None],
        ["صافي الرصيد (إجمالي الديون + الأرصدة الدائنة)", k["net"], None],
        ["متأخرات أكثر من 30 يوم", k["severe"], k["severe_share"]],
    ], columns=["البند", "المبلغ", "النسبة"])
    _write(ws, extra, start_row=nxt, money=("المبلغ",), pct=("النسبة",), title="معلومات إضافية")

    # 2. aging
    ws2 = _sheet(wb, "أعمار الديون")
    aging = r.aging.rename(columns={"Aging Bucket": "الفئة", "Amount Past Due": "المبلغ المتأخر", "% of Overdue": "% من المتأخرات",
                                    "% of Outstanding": "% من إجمالي الديون", "Invoices": "عدد الفواتير", "Clients": "عدد العملاء",
                                    "Avg Invoice": "متوسط الفاتورة", "Risk Level": "درجة الخطورة",
                                    "Recommended Action": "الإجراء المقترح"}).drop(columns=["key", "color"])
    tot = pd.DataFrame([["إجمالي المتأخرات", k["overdue"], 1.0 if k["overdue"] else None, k["overdue_ratio"], k["overdue_invoices"],
                         k["overdue_clients"], k["overdue"] / k["overdue_invoices"] if k["overdue_invoices"] else None, None, None]],
                       columns=aging.columns)
    nxt = _write(ws2, pd.concat([aging, tot], ignore_index=True), money=("المبلغ المتأخر", "متوسط الفاتورة"),
                 pct=("% من المتأخرات", "% من إجمالي الديون"), ints=("عدد الفواتير", "عدد العملاء"),
                 title="أعمار المتأخرات", fills={i: col for i, col in enumerate(r.aging["color"])}, total=True,
                 widths={"الإجراء المقترح": 48})
    nd = r.not_due.rename(columns={"Window": "موعد الاستحقاق", "Amount": "المبلغ", "Invoices": "عدد الفواتير"})
    _write(ws2, nd, start_row=nxt + 1, money=("المبلغ",), ints=("عدد الفواتير",), title="غير مستحقة بعد")

    # 3. weekly trend
    wk = pd.DataFrame([{
        "نهاية الأسبوع": x.end_date, "متأخرات بداية الأسبوع": x.start_overdue, "المحصّل منها": x.collected_overdue,
        "نسبة التحصيل": x.collection_rate, "إشعارات دائنة وتسويات": x.adjustments, "المتبقي منها": x.remaining,
        "متأخرات جديدة": x.new_overdue, "متأخرات نهاية الأسبوع": x.end_overdue, "المستهدف": x.target,
        "نسبة تحقيق المستهدف": x.achievement, "إجمالي التحصيل النقدي": x.collected_total} for x in r.weeks])
    _table_sheet(wb, "التحصيل الأسبوعي", wk, dates=("نهاية الأسبوع",), pct=("نسبة التحصيل", "نسبة تحقيق المستهدف"),
                 money=tuple(col for col in wk.columns if col not in ("نهاية الأسبوع", "نسبة التحصيل", "نسبة تحقيق المستهدف")))

    # 4. priority invoices
    pr = r.priority.rename(columns={"Client": "العميل", "Company": "الشركة", "Invoice": "رقم الفاتورة", "Due Date": "تاريخ الاستحقاق",
                                    "Outstanding": "المبلغ المستحق", "Days Overdue": "أيام التأخير", "Aging Bucket": "الفئة",
                                    "Recommended Action": "الإجراء المقترح"})
    _table_sheet(wb, "فواتير ذات أولوية", pr[["العميل", "الشركة", "رقم الفاتورة", "تاريخ الاستحقاق", "المبلغ المستحق", "أيام التأخير",
                                              "الفئة", "الإجراء المقترح"]],
                 money=("المبلغ المستحق",), dates=("تاريخ الاستحقاق",), ints=("أيام التأخير",), widths={"الإجراء المقترح": 60})

    # 5. clients
    cl = r.clients.rename(columns={"Client": "العميل", "Overdue": "المتأخر", "Not Due": "غير المستحق",
                                   "Unapplied Credit": "رصيد دائن غير مخصص", "Total Exposure": "إجمالي المديونية",
                                   "Invoices": "عدد الفواتير المتأخرة", "Oldest Days": "أقدم تأخير (يوم)",
                                   "Oldest Bucket": "أقدم فئة", "Flag": "ملاحظة"})
    _table_sheet(wb, "أكبر العملاء", cl[["العميل", "المتأخر", "غير المستحق", "رصيد دائن غير مخصص", "إجمالي المديونية",
                                       "عدد الفواتير المتأخرة", "أقدم تأخير (يوم)", "أقدم فئة", "ملاحظة"]],
                 money=("المتأخر", "غير المستحق", "رصيد دائن غير مخصص", "إجمالي المديونية"),
                 ints=("عدد الفواتير المتأخرة", "أقدم تأخير (يوم)"), widths={"ملاحظة": 50})

    # 6. deadlines and due next week
    cr = r.crossings.rename(columns={"Client": "العميل", "Invoice": "رقم الفاتورة", "Outstanding": "المبلغ المستحق",
                                     "Days Overdue": "أيام التأخير", "Moves To": "تنتقل إلى فئة", "On": "بتاريخ"})
    _table_sheet(wb, "مواعيد حرجة", cr, money=("المبلغ المستحق",), dates=("بتاريخ",), ints=("أيام التأخير",))
    up = (r.upcoming[["Client", "Invoice", "Due Date", "Outstanding"]] if len(r.upcoming)
          else pd.DataFrame(columns=["Client", "Invoice", "Due Date", "Outstanding"]))
    up = up.rename(columns={"Client": "العميل", "Invoice": "رقم الفاتورة", "Due Date": "تاريخ الاستحقاق", "Outstanding": "المبلغ المستحق"})
    _table_sheet(wb, "تستحق خلال 7 أيام", up, money=("المبلغ المستحق",), dates=("تاريخ الاستحقاق",))

    # 7. every open item
    oi = r.open_items.rename(columns={"Status": "الحالة", "customer": "العميل", "customer_ref": "رمز العميل", "document": "رقم المستند",
                                      "company": "الشركة", "invoice_date": "تاريخ الفاتورة", "due": "تاريخ الاستحقاق",
                                      "days": "أيام التأخير", "Aging Bucket": "الفئة", "res": "الرصيد المفتوح",
                                      "amount_document": "قيمة المستند", "currency": "العملة", "origin": "المرجع",
                                      "line_id": "رقم السطر في Odoo"})
    _table_sheet(wb, "كل الفواتير", oi[["الحالة", "العميل", "رمز العميل", "رقم المستند", "الشركة", "تاريخ الفاتورة", "تاريخ الاستحقاق",
                                       "أيام التأخير", "الفئة", "الرصيد المفتوح", "قيمة المستند", "العملة", "المرجع", "رقم السطر في Odoo"]],
                 money=("الرصيد المفتوح", "قيمة المستند"), dates=("تاريخ الفاتورة", "تاريخ الاستحقاق"), ints=("أيام التأخير",))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def team_workbook(summary: pd.DataFrame, path: Path) -> Path:
    wb = Workbook()
    ar = summary.rename(columns={
        "Salesperson": "المندوب", "Team": "الفريق", "Currency": "العملة", "Outstanding": "إجمالي الديون", "Overdue": "المتأخرات",
        "Not Due": "غير مستحقة", "Overdue Ratio": "نسبة المتأخرات", "31+ Days": "متأخرات أكثر من 30 يوم", "Debtors": "عدد العملاء",
        "Invoices": "عدد الفواتير", "Start Overdue": "متأخرات بداية الأسبوع", "Target": "المستهدف",
        "Collected (overdue)": "المحصّل منها", "Collection Rate": "نسبة التحصيل", "Target Achievement": "نسبة تحقيق المستهدف",
        "Adjustments": "إشعارات دائنة وتسويات", "Remaining From Start": "المتبقي منها", "New Overdue": "متأخرات جديدة",
        "End Overdue": "متأخرات نهاية الأسبوع", "Net Change": "صافي التغير", "Total Cash Collected": "إجمالي التحصيل النقدي",
        "Unapplied Credits": "أرصدة دائنة غير مخصصة"})
    ar["العملة"] = ar["العملة"].map(ccy_ar)
    pcts = ("نسبة المتأخرات", "نسبة التحصيل", "نسبة تحقيق المستهدف")
    ints = ("عدد العملاء", "عدد الفواتير")
    _table_sheet(wb, "ملخص الفريق", ar, pct=pcts, ints=ints,
                 money=tuple(col for col in ar.columns if col not in ("المندوب", "الفريق", "العملة") + pcts + ints))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
